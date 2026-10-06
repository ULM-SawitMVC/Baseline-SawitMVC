import docx
import os
import re
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_runs(paragraph, text):
    # Regex to capture **bold**, *italic*, `code`, and plain text
    # Match non-overlapping tokens
    pattern = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith('*') and token.endswith('*'):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        elif token.startswith('`') and token.endswith('`'):
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
        else:
            paragraph.add_run(token)

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    
    # 1 inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    for line in lines:
        raw_line = line.strip()
        if not raw_line:
            continue
        if raw_line == '---':
            continue

        if raw_line.startswith('# '):
            p = doc.add_paragraph(style='Heading 1')
            parse_runs(p, raw_line[2:])
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif raw_line.startswith('## '):
            p = doc.add_paragraph(style='Heading 2')
            parse_runs(p, raw_line[3:])
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif raw_line.startswith('### '):
            p = doc.add_paragraph(style='Heading 3')
            parse_runs(p, raw_line[4:])
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph(style='Normal')
            parse_runs(p, raw_line)
            p.paragraph_format.space_after = Pt(4)

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    here = os.path.dirname(os.path.abspath(__file__))
    convert_md_to_docx(
        os.path.join(here, 'Script.md'),
        os.path.join(here, 'Presentation_Script_ICIC2026_updated.docx')
    )
