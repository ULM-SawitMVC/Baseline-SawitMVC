#!/usr/bin/env python3
"""Convert the ICIC 2026 camera-ready LaTeX source into a two-column DOCX.

The source `main-new.tex` compiles to `main-new.pdf`, whose text is byte-identical
to the certified `2026318348.pdf`, so the DOCX carries the same content as the
submitted paper. Page geometry, paragraph styles and numbering come from the
official ICIC 2026 Word template stored in `template/`.
"""

from __future__ import annotations

import copy
import os
import re
import subprocess
import sys
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, Twips
from lxml import etree
from PIL import ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
TEX = os.path.join(HERE, "main-new.tex")
TEMPLATE = os.path.join(HERE, "template", "icic2026-conference-template-a4.docx")
FIGDIR = os.path.normpath(os.path.join(HERE, "..", "figures", "paper"))
OUT = os.path.join(HERE, "2026318348.docx")

# Template body section: A4 width 11906 twips, side margins 907, column gap 360;
# the title band uses side margins of 893.
TABLE_PT = 8.0
AUTHOR_PT = 9.0
COL_WIDTH_IN = (11906 - 2 * 907 - 360) / 2 / 1440
FULL_WIDTH_IN = (11906 - 2 * 907) / 1440
TITLE_WIDTH_IN = (11906 - 2 * 893) / 1440

ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]

SYMBOLS = {
    "alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4",
    "epsilon": "\u03b5", "varepsilon": "\u03b5", "theta": "\u03b8",
    "lambda": "\u03bb", "mu": "\u03bc", "sigma": "\u03c3", "tau": "\u03c4",
    "phi": "\u03c6", "Sigma": "\u03a3", "Delta": "\u0394", "Phi": "\u03a6",
    "pm": "\u00b1", "mp": "\u2213", "times": "\u00d7", "cdot": "\u00b7",
    "geq": "\u2265", "ge": "\u2265", "leq": "\u2264", "le": "\u2264",
    "neq": "\u2260", "approx": "\u2248", "equiv": "\u2261", "in": "\u2208",
    "notin": "\u2209", "subset": "\u2282", "mapsto": "\u21a6", "to": "\u2192",
    "rightarrow": "\u2192", "leftarrow": "\u2190", "infty": "\u221e",
    "sum": "\u2211", "prod": "\u220f", "int": "\u222b", "sqrt": "\u221a",
    "ldots": "\u2026", "cdots": "\u22ef", "dots": "\u2026",
    "circ": "\u2218", "ast": "*", "prime": "\u2032", "partial": "\u2202",
}

RELATIONS = set("=\u2208\u2265\u2264\u2260\u2248\u2261\u21a6\u2192<>")
BINARIES = set("+\u2212\u00d7\u00b7")
ROMAN_OPS = {"max", "min", "log", "exp", "sin", "cos", "clip", "det", "arg",
             "mathrm", "text", "operatorname", "mathsf", "mathit", "textrm"}


# --------------------------------------------------------------------------- #
# LaTeX reading helpers
# --------------------------------------------------------------------------- #

def strip_comments(src: str) -> str:
    out = []
    for line in src.split("\n"):
        buf, i = [], 0
        while i < len(line):
            if line[i] == "%" and (i == 0 or line[i - 1] != "\\"):
                break
            buf.append(line[i])
            i += 1
        out.append("".join(buf))
    return "\n".join(out)


def read_group(s: str, i: int):
    """Read a {...} group (or one character) starting at index i."""
    while i < len(s) and s[i] in " \t\r\n":
        i += 1                                  # TeX drops spaces after a control word
    if i >= len(s):
        return "", i
    if s[i] != "{":
        return s[i], i + 1
    depth, j = 0, i
    while j < len(s):
        if s[j] == "{" and (j == i or s[j - 1] != "\\"):
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


def split_top(s: str, sep: str):
    """Split on a separator that is not nested inside braces."""
    parts, buf, depth, i = [], [], 0, 0
    while i < len(s):
        c = s[i]
        if depth == 0 and s.startswith(sep, i):
            parts.append("".join(buf))
            buf = []
            i += len(sep)
            continue
        if c == "\\" and i + 1 < len(s) and not s[i + 1].isalpha():
            buf.append(s[i:i + 2])
            i += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
        buf.append(c)
        i += 1
    parts.append("".join(buf))
    return parts


# --------------------------------------------------------------------------- #
# Math mode -> styled runs
# --------------------------------------------------------------------------- #

class Tok:
    __slots__ = ("text", "style", "kind")

    def __init__(self, text, style, kind="ord"):
        self.text = text
        self.style = style
        self.kind = kind


def parse_math(src: str, base: dict | None = None, roman: bool = False):
    base = dict(base or {})
    toks: list[Tok] = []
    i, n = 0, len(src)

    def emit(text, extra=None, kind="ord"):
        if not text:
            return
        st = dict(base)
        if extra:
            st.update(extra)
        toks.append(Tok(text, st, kind))

    while i < n:
        ch = src[i]
        if ch == "\\":
            m = re.match(r"\\([A-Za-z]+)", src[i:])
            if not m:
                nxt = src[i + 1] if i + 1 < n else ""
                if nxt in "{}%$&_#":
                    emit(nxt)
                elif nxt == " ":
                    emit(" ")
                i += 2
                continue
            cmd, i = m.group(1), i + m.end()
            if cmd in ("left", "right", "big", "Big", "bigl", "bigr", "displaystyle"):
                continue
            if cmd in ("quad", "qquad"):
                emit(" ")
                continue
            if cmd in ("text", "mathrm", "operatorname", "textrm", "mathsf"):
                grp, i = read_group(src, i)
                toks.extend(parse_math(grp, base, roman=True))
                continue
            if cmd in ("mathbf", "boldsymbol", "bm"):
                grp, i = read_group(src, i)
                toks.extend(parse_math(grp, {**base, "bold": True}, roman))
                continue
            if cmd in ("mathit", "emph"):
                grp, i = read_group(src, i)
                toks.extend(parse_math(grp, {**base, "italic": True}, False))
                continue
            if cmd == "hat":
                grp, i = read_group(src, i)
                sub = parse_math(grp, base, roman)
                if sub:
                    sub[0].text = sub[0].text[:1] + "\u0302" + sub[0].text[1:]
                toks.extend(sub)
                continue
            if cmd in ("overline", "bar"):
                grp, i = read_group(src, i)
                sub = parse_math(grp, base, roman)
                plain = "".join(t.text for t in sub)
                # One letter takes a macron; a multi-letter name takes the
                # connecting overline on every character so the bar spans it.
                mark = "\u0304" if len(plain) == 1 else "\u0305"
                for t in sub:
                    t.text = "".join(ch + mark for ch in t.text)
                toks.extend(sub)
                continue
            if cmd == "frac":
                num, i = read_group(src, i)
                den, i = read_group(src, i)
                toks.extend(parse_math(num, base, roman))
                emit("/")
                dtoks = parse_math(den, base, roman)
                plain = "".join(t.text for t in dtoks)
                if len(plain) > 1:
                    emit("(")
                    toks.extend(dtoks)
                    emit(")")
                else:
                    toks.extend(dtoks)
                continue
            if cmd in ROMAN_OPS:
                emit(cmd, kind="op")
                continue
            if cmd in SYMBOLS:
                sym = SYMBOLS[cmd]
                kind = "rel" if sym in RELATIONS else ("bin" if sym in BINARIES else "ord")
                if cmd == "pm":
                    kind = "ord"
                if cmd in ("sum", "prod", "int"):
                    kind = "op"
                emit(sym, kind=kind)
                continue
            if cmd == "mathbf1":
                emit("1", {"bold": True})
                continue
            continue

        if ch in "_^":
            key = "sub" if ch == "_" else "sup"
            grp, j = read_group(src, i + 1)
            i = j
            for t in parse_math(grp, base, roman):
                st = dict(t.style)
                st[key] = True
                st.pop("spaced", None)
                toks.append(Tok(t.text, st, "script"))
            continue

        if ch == "{":
            grp, j = read_group(src, i)
            i = j
            inner = parse_math(grp, base, roman)
            if len(inner) == 1 and inner[0].text in BINARIES | RELATIONS:
                inner[0].kind = "ord"          # {+} means "tight"
            toks.extend(inner)
            continue

        if ch == "}":
            i += 1
            continue

        if ch == "-":
            emit("\u2212", kind="bin")
            i += 1
            continue

        if ch in RELATIONS:
            emit(ch, kind="rel")
            i += 1
            continue

        if ch in BINARIES:
            emit(ch, kind="bin")
            i += 1
            continue

        if ch == " ":
            i += 1
            continue

        if ch.isalpha() and not roman:
            emit(ch, {"italic": True})
            i += 1
            continue

        emit(ch)
        i += 1

    return toks


def space_math(toks: list[Tok]) -> list[Tok]:
    """Insert LaTeX-like spacing around relations, binaries and list commas."""
    out: list[Tok] = []
    after_op = False
    for idx, t in enumerate(toks):
        script = t.style.get("sub") or t.style.get("sup")
        if t.kind == "op" and out and (out[-1].text[-1:].isalnum() or out[-1].text[-1:] in ")]}"):
            out.append(Tok(" ", {}))
        if after_op and t.kind != "script" and not script:
            if t.text[:1].isalnum() or t.text[:1] in "(∑":
                out.append(Tok(" ", {}))
            after_op = False
        if t.kind == "op":
            after_op = True
        elif t.kind != "script":
            after_op = False
        prev = out[-1] if out else None
        prev_script = bool(prev and (prev.style.get("sub") or prev.style.get("sup")))
        if not script:
            if t.kind == "rel":
                if prev and not prev.text.endswith(" "):
                    out.append(Tok(" ", {}))
                out.append(t)
                out.append(Tok(" ", {}))
                continue
            if t.kind == "bin" and prev is not None and (
                    prev.text[-1:].isalnum() or prev.text[-1:] in ")]}\u0304"):
                out.append(Tok(" ", {}))
                out.append(t)
                out.append(Tok(" ", {}))
                continue
            if t.text == ",":
                out.append(t)
                out.append(Tok(" ", {}))
                continue
        out.append(t)
    # collapse duplicate spaces
    merged: list[Tok] = []
    for t in out:
        if merged and t.text == " " and merged[-1].text.endswith(" "):
            continue
        merged.append(t)
    while merged and merged[-1].text == " ":
        merged.pop()
    return merged


# --------------------------------------------------------------------------- #
# Text mode -> styled runs
# --------------------------------------------------------------------------- #

class Converter:
    def __init__(self, bibkeys, labels):
        self.bib = {k: i + 1 for i, k in enumerate(bibkeys)}
        self.labels = labels
        self.footnotes: list[str] = []

    def cite(self, keys):
        nums = sorted({self.bib[k.strip()] for k in keys.split(",") if k.strip() in self.bib})
        groups, run = [], [nums[0]] if nums else []
        for v in nums[1:]:
            if v == run[-1] + 1:
                run.append(v)
            else:
                groups.append(run)
                run = [v]
        if run:
            groups.append(run)
        parts = []
        for g in groups:
            if len(g) >= 3:
                parts.append("[%d]\u2013[%d]" % (g[0], g[-1]))
            else:
                parts.extend("[%d]" % v for v in g)
        return ", ".join(parts)

    def inline(self, text: str, base: dict | None = None) -> list[Tok]:
        base = dict(base or {})
        toks: list[Tok] = []
        i, n = 0, len(text)

        def emit(s, extra=None):
            if not s:
                return
            st = dict(base)
            if extra:
                st.update(extra)
            toks.append(Tok(s, st))

        while i < n:
            ch = text[i]
            if ch == "$":
                j = text.index("$", i + 1)
                mtoks = space_math(parse_math(text[i + 1:j]))
                for t in mtoks:
                    st = dict(base)
                    st.update(t.style)
                    toks.append(Tok(t.text, st))
                i = j + 1
                continue
            if ch == "\\":
                m = re.match(r"\\([A-Za-z]+)\*?", text[i:])
                if not m:
                    nxt = text[i + 1] if i + 1 < n else ""
                    if nxt in "%&_#${}":
                        emit(nxt)
                    elif nxt == " ":
                        emit(" ")
                    i += 2
                    continue
                cmd, i = m.group(1), i + m.end()
                if cmd in ("textbf", "bf"):
                    grp, i = read_group(text, i)
                    toks.extend(self.inline(grp, {**base, "bold": True}))
                elif cmd in ("textit", "emph", "it"):
                    grp, i = read_group(text, i)
                    toks.extend(self.inline(grp, {**base, "italic": True}))
                elif cmd == "textsuperscript" or cmd == "affmark":
                    grp, i = read_group(text, i)
                    toks.extend(self.inline(grp, {**base, "sup": True}))
                elif cmd == "url":
                    grp, i = read_group(text, i)
                    emit(grp)
                elif cmd == "cite":
                    grp, i = read_group(text, i)
                    emit(self.cite(grp))
                elif cmd == "ref":
                    grp, i = read_group(text, i)
                    emit(self.labels.get(grp.strip(), "?"))
                elif cmd == "footnote":
                    grp, i = read_group(text, i)
                    self.footnotes.append("".join(t.text for t in self.inline(grp)))
                    emit(str(len(self.footnotes)), {"sup": True, "footnote": len(self.footnotes)})
                elif cmd == "footnotesize":
                    pass
                elif cmd in ("label", "vspace", "hspace", "index"):
                    _, i = read_group(text, i)
                elif cmd in ("par", "raggedright", "centering", "balance",
                             "noindent", "smallskip", "medskip", "bigskip"):
                    pass
                else:
                    pass
                continue
            if text.startswith("---", i):
                emit("\u2014")
                i += 3
                continue
            if text.startswith("--", i):
                emit("\u2013")
                i += 2
                continue
            if text.startswith("``", i):
                emit("\u201c")
                i += 2
                continue
            if text.startswith("''", i):
                emit("\u201d")
                i += 2
                continue
            if ch == "~":
                emit("\u00a0")
                i += 1
                continue
            if ch in "{}":
                i += 1
                continue
            emit(ch)
            i += 1

        return toks


# --------------------------------------------------------------------------- #
# DOCX helpers
# --------------------------------------------------------------------------- #

def set_font(run, size=None, bold=False, italic=False, name="Times New Roman"):
    """Direct run formatting; size None keeps the paragraph style's size."""
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


def add_runs(par, toks, size, base_bold=False, base_italic=False, smallcaps=False):
    for t in toks:
        if not t.text:
            continue
        if t.style.get("footnote"):
            run = par.add_run()
            run.font.superscript = True
            ref = OxmlElement("w:footnoteReference")
            ref.set(qn("w:id"), str(t.style["footnote"]))
            run._r.append(ref)
            continue
        run = par.add_run(t.text)
        set_font(run, size,
                 bold=t.style.get("bold", base_bold),
                 italic=t.style.get("italic", base_italic))
        if t.style.get("sub"):
            run.font.subscript = True
        if t.style.get("sup"):
            run.font.superscript = True
        if smallcaps:
            rpr = run._element.get_or_add_rPr()
            el = OxmlElement("w:smallCaps")
            el.set(qn("w:val"), "1")
            rpr.append(el)
    return par


def use_style(par, name, *, align=None, before=None, after=None, first=None,
              left=None, keep=None, line=None, numbered=True):
    """Apply a template paragraph style; keyword arguments (points) override it.

    numbered=False detaches the style's automatic list numbering, for captions
    that live in text boxes, where Word does not continue the main sequence.
    """
    par.style = name
    pf = par.paragraph_format
    if align is not None:
        par.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if first is not None:
        pf.first_line_indent = Pt(first)
    if left is not None:
        pf.left_indent = Pt(left)
    if keep is not None:
        pf.keep_with_next = keep
    if line is not None:
        pf.line_spacing = line
    if not numbered:
        num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = 0
    return par


def set_cell_border(cell, edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in edges:
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(edges[edge]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
            borders.append(el)
    tc_pr.append(borders)


def set_cell_margins(cell, top=10, bottom=10, left=22, right=22):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + name)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


_TABLE_FONTS = {}


def text_width_in(text, bold=True, size=TABLE_PT):
    """Printed width in inches of text in Times New Roman, from Windows font metrics."""
    key = (bold, size)
    if key not in _TABLE_FONTS:
        name = "timesbd.ttf" if bold else "times.ttf"
        path = os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", name)
        _TABLE_FONTS[key] = ImageFont.truetype(path, int(size * 10))   # 10x for precision
    return _TABLE_FONTS[key].getlength(text) / 10 / 72


def column_fractions(rows, ncols, widths, total_in):
    """Column width fractions: explicit p{} widths win, otherwise content-based.

    Content-based widths are then raised so that no column is narrower than its
    longest single word, which is what makes headings such as "Wrong" break.
    """
    lens = [[] for _ in range(ncols)]
    words = [[] for _ in range(ncols)]
    for row in rows:
        ci = 0
        for span, _al, content in row:
            if span == 1 and ci < ncols:
                plain = re.sub(r"\\[A-Za-z]+", " ", content)
                plain = re.sub(r"[${}\\]", "", plain).strip()
                lens[ci].append(len(plain))
                words[ci].extend(plain.split())
            ci += span

    if any(w for w in widths):
        known = sum(w for w in widths if w)
        free = [i for i, w in enumerate(widths) if not w]
        rest = max(0.05, (1.0 - known) / len(free)) if free else 0.0
        fracs = [w if w else rest for w in widths]
    else:
        weights = [max(3.0, max(c) if c else 3.0) ** 0.9 for c in lens]
        fracs = [w / sum(weights) for w in weights]

    # Narrowest usable width per column: its widest word measured in bold Times
    # New Roman, plus cell margins. A character-count estimate broke "mAP50".
    mins = [(max((text_width_in(w) for w in ws), default=0.1) + 0.06) / total_in for ws in words]
    if sum(mins) >= 0.98:                       # every word cannot fit; share by need
        return [m / sum(mins) for m in mins]
    for _ in range(ncols):
        short = [i for i in range(ncols) if fracs[i] < mins[i]]
        if not short:
            break
        fixed = sum(mins[i] for i in short)
        free = [i for i in range(ncols) if i not in short]
        pool = sum(fracs[i] for i in free)
        scale = (1.0 - fixed) / pool if pool else 0.0
        fracs = [mins[i] if i in short else fracs[i] * scale for i in range(ncols)]
    total = sum(fracs)
    return [f / total for f in fracs]


def set_table_widths(table, fracs, total_in):
    """Fixed layout with per-column widths, honouring horizontally merged cells."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.get_or_add_tblLayout()
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(total_in * 1440)))

    twips = [int(total_in * f * 1440) for f in fracs]
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, tw in zip(grid.findall(qn("w:gridCol")), twips):
            col.set(qn("w:w"), str(tw))

    for tr in table._tbl.findall(qn("w:tr")):
        ci = 0
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            span = 1
            if tc_pr is not None:
                gs = tc_pr.find(qn("w:gridSpan"))
                if gs is not None:
                    span = int(gs.get(qn("w:val")))
            width = sum(twips[ci:ci + span])
            if tc_pr is None:
                tc_pr = OxmlElement("w:tcPr")
                tc.insert(0, tc_pr)
            tcw = tc_pr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tc_pr.insert(0, tcw)
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(width))
            ci += span


# --------------------------------------------------------------------------- #
# Structure extraction
# --------------------------------------------------------------------------- #

def build_labels(body: str):
    labels, sec, sub, tab, fig = {}, 0, 0, 0, 0
    env = None
    for line in body.split("\n"):
        s = line.strip()
        if s.startswith("\\section*"):
            env = None
            continue
        if s.startswith("\\section{"):
            sec += 1
            sub = 0
            env = None
            continue
        if s.startswith("\\subsection{"):
            sub += 1
            env = None
            continue
        if re.match(r"\\begin\{table\*?\}", s):
            tab += 1
            env = "table"
            continue
        if re.match(r"\\begin\{figure\*?\}", s):
            fig += 1
            env = "figure"
            continue
        if re.match(r"\\end\{(table|figure)\*?\}", s):
            env = None
            continue
        m = re.search(r"\\label\{([^}]+)\}", s)
        if m:
            key = m.group(1)
            if env == "table":
                labels[key] = ROMAN[tab]
            elif env == "figure":
                labels[key] = str(fig)
            else:
                labels[key] = ROMAN[sec] + ("-" + chr(64 + sub) if sub else "")
    return labels


def parse_tabular(chunk: str):
    """Return (colspec, rows) where each row is a list of (span, align, cell)."""
    m = re.search(r"\\begin\{tabular\*?\}", chunk)
    i = m.end()
    if m.group(0).endswith("*}") and chunk[i] == "{":   # tabular* width argument
        _, i = read_group(chunk, i)
    spec, i = read_group(chunk, i)
    end = chunk.index("\\end{tabular", i)
    body = chunk[i:end]

    spec_clean = re.sub(r"@\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", "", spec)
    spec_clean = re.sub(r">\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", "", spec_clean)
    aligns, widths = [], []
    k = 0
    while k < len(spec_clean):
        c = spec_clean[k]
        if c in "lcr":
            aligns.append(c)
            widths.append(None)
            k += 1
        elif c == "p":
            grp, k = read_group(spec_clean, k + 1)
            aligns.append("l")
            mw = re.search(r"([0-9.]+)\\?(?:columnwidth|textwidth)", grp)
            widths.append(float(mw.group(1)) if mw else None)
        else:
            k += 1

    rows = []
    for raw in split_top(body, "\\\\"):
        raw = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", raw)
        raw = re.sub(r"\\(toprule|midrule|bottomrule)\b", "", raw)
        raw = re.sub(r"\\addlinespace(\[[^\]]*\])?", "", raw)
        raw = re.sub(r"^\[[0-9.]+pt\]", "", raw.strip())
        if not raw.strip():
            continue
        cells = []
        for cell in split_top(raw, "&"):
            cell = cell.strip()
            if cell.startswith("\\multicolumn"):
                k = len("\\multicolumn")
                span_s, k = read_group(cell, k)
                spec_s, k = read_group(cell, k)
                content, _ = read_group(cell, k)
                span = int(span_s)
                letters = re.sub(r"@\{[^{}]*\}", "", spec_s)
                al = "l" if "l" in letters else ("r" if "r" in letters else "c")
                cells.append((span, al, content.strip()))
            else:
                cells.append((1, None, cell))
        rows.append(cells)
    return aligns, widths, rows


def rule_rows(chunk: str):
    """Row indices (0-based, counting content rows) that carry booktabs rules."""
    m = re.search(r"\\begin\{tabular\*?\}", chunk)
    i = m.end()
    if m.group(0).endswith("*}") and chunk[i] == "{":
        _, i = read_group(chunk, i)
    _, i = read_group(chunk, i)
    body = chunk[i:chunk.index("\\end{tabular", i)]
    mids, idx = [], 0
    for raw in split_top(body, "\\\\"):
        has_mid = "\\midrule" in raw
        cmids = re.findall(r"\\cmidrule(?:\([^)]*\))?\{(\d+)-(\d+)\}", raw)
        stripped = re.sub(r"\\(toprule|midrule|bottomrule)\b", "", raw)
        stripped = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", stripped)
        stripped = re.sub(r"\\addlinespace(\[[^\]]*\])?", "", stripped)
        stripped = re.sub(r"^\[[0-9.]+pt\]", "", stripped.strip())
        if not stripped.strip():
            continue
        if has_mid:
            mids.append(("mid", idx))
        for a, b in cmids:
            mids.append(("cmid", idx, int(a), int(b)))
        idx += 1
    return mids, idx


# --------------------------------------------------------------------------- #
# Document builder
# --------------------------------------------------------------------------- #

def write_footnotes(doc, notes, unnumbered=()):
    """Append real footnotes to the template's footnotes part.

    python-docx has no footnote API, but the template already carries the
    part with its separator entries, so only the notes themselves are added.
    Notes whose 1-based id is in `unnumbered` carry no mark, as IEEE asks for
    the first-page sponsor footnote.
    """
    if not notes:
        return
    part = next(rel.target_part for rel in doc.part.rels.values()
                if rel.reltype.endswith("/footnotes"))
    root = getattr(part, "_element", None)
    is_xml_part = root is not None
    if not is_xml_part:
        root = parse_xml(part.blob)
    for n, note in enumerate(notes, 1):
        numbered = n not in unnumbered
        mark = ('<w:r><w:rPr><w:vertAlign w:val="superscript"/><w:sz w:val="16"/></w:rPr>'
                '<w:footnoteRef/></w:r>') if numbered else ''
        root.append(parse_xml(
            '<w:footnote %s w:id="%d"><w:p><w:pPr><w:jc w:val="both"/></w:pPr>%s'
            '<w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r>'
            '</w:p></w:footnote>'
            % (nsdecls("w"), n, mark, escape((" " if numbered else "") + note))))
    if not is_xml_part:
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                    standalone=True)


def main():
    src = strip_comments(open(TEX, encoding="utf-8").read())
    body = src[src.index("\\begin{document}"):]

    bibkeys = re.findall(r"\\bibitem\{([^}]+)\}", src)
    labels = build_labels(body)
    conv = Converter(bibkeys, labels)

    # The template supplies page geometry, paragraph styles and numbering; its
    # sample content is discarded. Of its five sections, 0 (title band) and
    # 3 (two-column body) are reused as they are.
    doc = Document(TEMPLATE)
    body_el = doc.element.body
    sect_prs = [copy.deepcopy(s) for s in body_el.iter(qn("w:sectPr"))]
    for child in list(body_el):
        body_el.remove(child)
    title_sect, body_sect = sect_prs[0], sect_prs[3]
    for ref in title_sect.findall(qn("w:footerReference")):   # copyright placeholder
        title_sect.remove(ref)
        doc.part.drop_rel(ref.get(qn("r:id")))
    for el in title_sect.findall(qn("w:titlePg")):
        title_sect.remove(el)
    body_el.append(body_sect)

    state = {"sec": 0, "sub": 0, "tab": 0, "fig": 0, "eq": 0, "float": 0}

    # ---------------- title band: title and author blocks ----------------
    title_src, _ = read_group(body, body.index("\\title{") + len("\\title"))
    thanks = ""
    thanks_m = re.search(r"\\thanks\{", title_src)
    if thanks_m:
        thanks, _ = read_group(title_src, thanks_m.end() - 1)
        title_src = title_src[:thanks_m.start()]
    title_txt = " ".join(title_src.split())
    p = doc.add_paragraph()
    use_style(p, "paper title")
    add_runs(p, conv.inline(title_txt), None)
    unnumbered = set()
    if thanks:
        # Sponsor acknowledgments go in an unnumbered first-page footnote: a
        # footnote reference whose custom mark is empty.
        conv.footnotes.append(" ".join("".join(t.text for t in conv.inline(thanks)).split()))
        unnumbered.add(len(conv.footnotes))
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:customMarkFollows"), "1")
        ref.set(qn("w:id"), str(len(conv.footnotes)))
        # Referenced from the first body paragraph, so the note sits under the
        # first column rather than across the title band.
        state["thanks_ref"] = ref

    # One block per author (name, department, organization, city, email), left
    # to right; \linebreakand starts a new row. A borderless table keeps the
    # blocks of a row top-aligned, which column breaks inside a section do not.
    author_src, _ = read_group(body, body.index("\\author{") + len("\\author"))
    rows = [[]]
    for part in re.split(r"(\\linebreakand\b|\\and\b)", author_src):
        if part == "\\linebreakand":
            rows.append([])
        elif part.strip() and part != "\\and":
            name, _ = read_group(part, re.search(r"\\IEEEauthorblockN\{", part).end() - 1)
            aff, _ = read_group(part, re.search(r"\\IEEEauthorblockA\{", part).end() - 1)
            lines = [" ".join(x.split()) for x in split_top(aff, "\\\\")]
            rows[-1].append((" ".join(name.split()), [x for x in lines if x]))
    ncols = max(len(row) for row in rows)
    grid = doc.add_table(rows=len(rows), cols=ncols)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, row in enumerate(rows):
        for k, (name, lines) in enumerate(row):
            cell = grid.cell(r_i, k)
            set_cell_margins(cell, top=0, bottom=0, left=30, right=30)
            par = cell.paragraphs[0]
            use_style(par, "Author", before=None if r_i == 0 else 10, after=0)
            add_runs(par, conv.inline(name), AUTHOR_PT)
            for line in lines:
                par.add_run().add_break()
                add_runs(par, conv.inline(line), AUTHOR_PT)
    set_table_widths(grid, [1 / ncols] * ncols, TITLE_WIDTH_IN)
    end = doc.add_paragraph()
    use_style(end, "Normal", before=0, after=0, line=Pt(12))
    end._p.get_or_add_pPr().append(title_sect)

    props = doc.core_properties
    props.title = title_txt
    props.author = ", ".join(name for row in rows for name, _ in row)
    props.last_modified_by = ""
    props.subject = props.keywords = props.comments = ""

    # ---------------- abstract & keywords ----------------
    abst = re.search(r"\\begin\{abstract\}(.+?)\\end\{abstract\}", body, re.S).group(1)
    p = doc.add_paragraph()
    use_style(p, "Abstract")
    r = p.add_run("Abstract. ")
    set_font(r, bold=True, italic=True)
    add_runs(p, conv.inline(" ".join(abst.split())), None, base_bold=True)

    kw = re.search(r"\\begin\{IEEEkeywords\}(.+?)\\end\{IEEEkeywords\}", body, re.S).group(1)
    p = doc.add_paragraph()
    use_style(p, "Keywords")
    r = p.add_run("Keywords. ")
    set_font(r, bold=True, italic=True)
    add_runs(p, conv.inline(" ".join(kw.split())), None, base_bold=True, base_italic=True)

    def add_body_paragraph(text):
        text = " ".join(text.split())
        if text:
            p = doc.add_paragraph()
            use_style(p, "Body Text")
            add_runs(p, conv.inline(text), None)
            if state.get("thanks_ref") is not None:
                p.add_run()._r.append(state.pop("thanks_ref"))

    # ---- floats ----
    def new_anchor():
        anchor = doc.add_paragraph()
        use_style(anchor, "Normal", before=0, after=0, line=Pt(1))
        return anchor

    def add_float(anchor, blocks, width_in, height_emu, wide):
        """Put blocks in a floating text box: page top (wide) or its column.

        IEEE places figures and tables at column tops and wide figures across
        the page top. A continuous section break cannot do this: Word balances
        the columns above such a break, splitting the paragraph before a wide
        figure across both columns. A text box with top-and-bottom wrapping
        leaves the column flow untouched. Word grows the box to fit its content
        (spAutoFit), so height_emu is only a first estimate.

        Column boxes start at their anchor paragraph. place_floats.ps1 then
        moves each box to a column top inside Word, because only Word's
        layout knows the page and column an anchor lands in; it also
        re-anchors a box whose page would otherwise alternate.
        """
        state["float"] += 1
        # Drawing ids must be unique. A fixed numbering scheme collided with
        # the ids python-docx gives pictures created after an earlier float.
        uid = doc.part.next_id
        if wide:
            vertical = '<wp:positionV relativeFrom="margin"><wp:align>top</wp:align></wp:positionV>'
        else:
            vertical = ('<wp:positionV relativeFrom="paragraph">'
                        '<wp:posOffset>0</wp:posOffset></wp:positionV>')
        xml = (
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            '<mc:Choice xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" Requires="wps">'
            '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<wp:anchor distT="0" distB="{db}" distL="0" distR="0" simplePos="0"'
            ' relativeHeight="{uid}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="0">'
            '<wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="{hfrom}"><wp:align>center</wp:align></wp:positionH>'
            '{vertical}'
            '<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            '<wp:wrapTopAndBottom/>'
            '<wp:docPr id="{uid}" name="{kind} {n}"/><wp:cNvGraphicFramePr/>'
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '<wps:wsp><wps:cNvSpPr txBox="1"/>'
            '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
            '<wps:txbx><w:txbxContent/></wps:txbx>'
            '<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0"'
            ' anchor="t" anchorCtr="0"><a:spAutoFit/></wps:bodyPr>'
            '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'
            '</mc:Choice></mc:AlternateContent>'
        ).format(db=int(Pt(10)), uid=uid, n=state["float"], kind="Wide" if wide else "Column",
                 cx=int(Inches(width_in)), cy=int(height_emu),
                 hfrom="margin" if wide else "column", vertical=vertical)
        alt = parse_xml(xml)
        content = alt.find(".//" + qn("w:txbxContent"))
        for el in blocks:
            content.append(el)
        run = anchor.add_run()
        set_font(run, 1)
        run._r.append(alt)

    # ---- figures ----
    def add_figure(chunk, wide):
        state["fig"] += 1
        src_m = re.search(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", chunk)
        path = os.path.join(FIGDIR, os.path.basename(src_m.group(2)))
        png = os.path.splitext(path)[0] + ".png"
        img = png if os.path.exists(png) else path
        cap_m = re.search(r"\\caption\{", chunk)
        caption, _ = read_group(chunk, cap_m.end() - 1)
        caption = " ".join(caption.split())
        width_in = FULL_WIDTH_IN if wide else COL_WIDTH_IN

        anchor = new_anchor()
        pic = doc.add_paragraph()
        use_style(pic, "Normal", before=0, after=0)
        shape = pic.add_run().add_picture(img, width=Inches(width_in))
        cap = doc.add_paragraph()
        use_style(cap, "figure caption", numbered=False, after=0)
        r = cap.add_run("Fig. %d. " % state["fig"])
        set_font(r)
        add_runs(cap, conv.inline(caption), None)

        lines = -(-(len(caption) + 8) // int(width_in * 20))   # ~20 chars/inch at 8 pt
        height = shape.height + Pt(4 + 10 * lines)
        add_float(anchor, [pic._p, cap._p], width_in, height, wide)

    # ---- tables ----
    def add_table(chunk):
        state["tab"] += 1
        num = ROMAN[state["tab"]]
        cap_m = re.search(r"\\caption\{", chunk)
        caption, _ = read_group(chunk, cap_m.end() - 1)
        caption = " ".join(caption.split())

        anchor = new_anchor()
        head = doc.add_paragraph()
        use_style(head, "table head", numbered=False, before=0, keep=True)
        r = head.add_run("TABLE %s. " % num)
        set_font(r)
        add_runs(head, conv.inline(caption), None)

        aligns, widths, rows = parse_tabular(chunk)
        rules, _ = rule_rows(chunk)
        ncols = max(sum(c[0] for c in row) for row in rows)
        while len(aligns) < ncols:
            aligns.append("c")
        while len(widths) < ncols:
            widths.append(None)
        header_rows = min((rule[1] for rule in rules if rule[0] == "mid"), default=0)

        table = doc.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        for ri, row in enumerate(rows):
            is_head = ri < header_rows
            ci = 0
            for span, al, content in row:
                cell = table.cell(ri, ci)
                if span > 1:
                    cell = cell.merge(table.cell(ri, ci + span - 1))
                set_cell_margins(cell)
                par = cell.paragraphs[0]
                use_style(par, "table col head" if is_head else "table copy",
                          align={"l": WD_ALIGN_PARAGRAPH.LEFT,
                                 "c": WD_ALIGN_PARAGRAPH.CENTER,
                                 "r": WD_ALIGN_PARAGRAPH.RIGHT}[al or aligns[ci]],
                          before=1, after=1)
                # Cells written over several indented source lines would keep
                # that indentation as runs of spaces.
                add_runs(par, conv.inline(" ".join(content.split())), None, base_bold=is_head)
                ci += span
            for c in range(ci, ncols):
                set_cell_margins(table.cell(ri, c))

        fracs = column_fractions(rows, ncols, widths, COL_WIDTH_IN)
        set_table_widths(table, fracs, COL_WIDTH_IN)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))

        last = len(rows) - 1
        for c in range(ncols):
            set_cell_border(table.cell(0, c), {"top": 12})
            set_cell_border(table.cell(last, c), {"bottom": 12})
        for rule in rules:
            if rule[0] == "mid":
                ri = rule[1] - 1
                if ri >= 0:
                    for c in range(ncols):
                        cur = {"bottom": 6}
                        if ri == 0:
                            cur["top"] = 12
                        set_cell_border(table.cell(ri, c), cur)
            else:
                _, ri, a, b = rule
                ri -= 1
                if ri >= 0:
                    for c in range(a - 1, min(b, ncols)):
                        cur = {"bottom": 6}
                        if ri == 0:
                            cur["top"] = 12
                        set_cell_border(table.cell(ri, c), cur)

        note_m = re.search(r"\\begin\{minipage\}.*?\\end\{minipage\}", chunk, re.S)
        notes = []
        if note_m:
            note = note_m.group(0)
            note = re.sub(r"\\begin\{minipage\}\{[^}]*\}", "", note)
            note = note.replace("\\end{minipage}", "")
            note = note.replace("\\footnotesize", "").replace("\\raggedright", "")
            # Lettered table footnotes are separated by \\ in the source.
            notes = [" ".join(x.split()) for x in split_top(note, "\\\\") if x.strip()]
        note_paras = []
        for k, text in enumerate(notes or [""]):   # a text box must end with a paragraph
            par = doc.add_paragraph()
            use_style(par, "table footnote", numbered=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      left=0, first=0, before=3 if k == 0 and text else 0, after=0)
            if text:
                add_runs(par, conv.inline(text), None)
            note_paras.append(par)

        per_line = int(COL_WIDTH_IN * 20)
        head_lines = -(-(len(caption) + 10) // per_line)
        note_lines = sum(-(-len(x) // int(COL_WIDTH_IN * 27)) for x in notes) or 1
        height = Pt(10 * head_lines + 6 + 12 * len(rows) + 7.5 * note_lines + 3)
        add_float(anchor, [head._p, table._tbl] + [par._p for par in note_paras],
                  COL_WIDTH_IN, height, False)

    # ---- equations ----
    def add_align(chunk):
        inner = re.search(r"\\begin\{align\}(.+?)\\end\{align\}", chunk, re.S).group(1)
        for line in split_top(inner, "\\\\"):
            line = line.strip()
            if not line:
                continue
            state["eq"] += 1
            p = doc.add_paragraph()
            use_style(p, "equation")
            tabs = p.paragraph_format.tab_stops
            for pos in (2520, 5040):          # template stops assume a wider column
                tabs.add_tab_stop(Twips(pos), WD_TAB_ALIGNMENT.CLEAR)
            tabs.add_tab_stop(Inches(COL_WIDTH_IN / 2), WD_TAB_ALIGNMENT.CENTER)
            tabs.add_tab_stop(Inches(COL_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run("\t")
            set_font(r)
            add_runs(p, space_math(parse_math(line.replace("&", ""))), None)
            r = p.add_run("\t(%d)" % state["eq"])
            set_font(r)

    # ---- walk the document body ----
    lines = body.split("\n")
    i = 0
    para: list[str] = []

    def flush_para():
        nonlocal para
        if para:
            add_body_paragraph(" ".join(para))
        para = []

    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if s.startswith("\\begin{document}") or s.startswith("\\maketitle"):
            i += 1
            continue
        if s.startswith("\\title{") or s.startswith("\\author{"):
            depth = 0
            while i < len(lines):
                depth += lines[i].count("{") - lines[i].count("}")
                i += 1
                if depth <= 0:
                    break
            continue
        if re.match(r"\\begin\{(abstract|IEEEkeywords)\}", s):
            while i < len(lines) and not re.match(r"\\end\{(abstract|IEEEkeywords)\}", lines[i].strip()):
                i += 1
            i += 1
            continue
        if s.startswith("\\end{document}"):
            break

        m = re.match(r"\\section\*\{(.+)\}$", s)
        if m:
            flush_para()
            p = doc.add_paragraph()
            use_style(p, "Heading 5")                  # unnumbered component head
            add_runs(p, conv.inline(m.group(1)), None)
            i += 1
            continue

        m = re.match(r"\\section\{(.+)\}$", s)
        if m:
            flush_para()
            state["sec"] += 1
            state["sub"] = 0
            p = doc.add_paragraph()
            use_style(p, "Heading 1")                  # template numbers it I., II., ...
            add_runs(p, conv.inline(m.group(1)), None)
            i += 1
            continue

        m = re.match(r"\\subsection\{(.+)\}$", s)
        if m:
            flush_para()
            state["sub"] += 1
            p = doc.add_paragraph()
            use_style(p, "Heading 2")                  # template numbers it A., B., ...
            add_runs(p, conv.inline(m.group(1)), None, base_italic=True)
            i += 1
            continue

        m = re.match(r"\\begin\{(table|figure|align)(\*?)\}", s)
        if m:
            flush_para()
            env, star = m.group(1), m.group(2)
            end_tag = "\\end{%s%s}" % (env, star)
            chunk_lines = [line]
            i += 1
            while i < len(lines) and end_tag not in lines[i]:
                chunk_lines.append(lines[i])
                i += 1
            chunk_lines.append(lines[i] if i < len(lines) else end_tag)
            i += 1
            chunk = "\n".join(chunk_lines)
            if env == "table":
                add_table(chunk)
            elif env == "figure":
                add_figure(chunk, wide=bool(star))
            else:
                add_align(chunk)
            continue

        if s.startswith("\\begin{thebibliography}"):
            flush_para()
            break

        if not s:
            flush_para()
            i += 1
            continue

        if s.startswith("\\label{") or s in ("\\balance", "\\centering"):
            i += 1
            continue

        para.append(s)
        i += 1

    flush_para()

    # ---- references ----
    p = doc.add_paragraph()
    use_style(p, "Heading 5")
    add_runs(p, conv.inline("References"), None)

    bib = src[src.index("\\begin{thebibliography}"):src.index("\\end{thebibliography}")]
    entries = re.split(r"\\bibitem\{[^}]+\}", bib)[1:]
    for entry in entries:
        p = doc.add_paragraph()
        use_style(p, "references")                     # template numbers it [1], [2], ...
        add_runs(p, conv.inline(" ".join(entry.split())), None)

    write_footnotes(doc, conv.footnotes, unnumbered)
    doc.save(OUT)
    print("wrote", OUT)
    subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
                    "-File", os.path.join(HERE, "place_floats.ps1"), "-Docx", OUT],
                   check=True)
    print("headings=%d tables=%d figures=%d equations=%d refs=%d footnotes=%d floats=%d"
          % (state["sec"], state["tab"], state["fig"], state["eq"], len(entries),
             len(conv.footnotes), state["float"]))


if __name__ == "__main__":
    sys.exit(main())
